import urllib.request
import sys
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from selenium import webdriver
from selenium.webdriver.common.by import By
import threading
import time
import os
import re
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ================= HỆ THỐNG AUTO-UPDATE =================
APP_VERSION = 2.0  # Mỗi lần đăng Github, bạn nhớ nâng số này lên (VD: 1.1)

VERSION_URL = "https://raw.githubusercontent.com/canxi99/Tool-Tai-Truyen/refs/heads/main/version.txt"
CODE_URL = "https://raw.githubusercontent.com/canxi99/Tool-Tai-Truyen/refs/heads/main/taitruyen.py"

def check_for_updates():
    try:
        req = urllib.request.Request(VERSION_URL, headers={'User-Agent': 'Mozilla/5.0'})
        response = urllib.request.urlopen(req, timeout=3)
        latest_version = float(response.read().decode('utf-8').strip())

        if latest_version > APP_VERSION:
            root_upd = tk.Tk()
            root_upd.withdraw()
            result = messagebox.askyesno("Có bản cập nhật mới!", 
                                         f"Tuyệt vời! Tác giả vừa ra mắt phiên bản v{latest_version}.\n"
                                         "Bạn có muốn phần mềm tự động tải và cập nhật ngay không?")
            if result:
                req_code = urllib.request.Request(CODE_URL, headers={'User-Agent': 'Mozilla/5.0'})
                new_code = urllib.request.urlopen(req_code, timeout=10).read().decode('utf-8')
                
                with open(sys.argv[0], 'w', encoding='utf-8') as f:
                    f.write(new_code)
                
                messagebox.showinfo("Thành công", "Đã cập nhật xong! Phần mềm sẽ tự khởi động lại.")
                os.execv(sys.executable, ['python'] + sys.argv)
            root_upd.destroy()
    except Exception:
        pass # Không có mạng hoặc link lỗi thì bỏ qua, cho xài bản cũ

# Chạy kiểm tra update ngay khi mở phần mềm
check_for_updates()

# ================= CÁC BIẾN TOÀN CỤC & CẤU HÌNH =================
CONFIG_FILE = "novel_configs.json"
pause_event = threading.Event()
pause_event.set()

def load_configs():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception: return {}
    return {}

def save_configs_to_file(configs):
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(configs, f, indent=4, ensure_ascii=False)

def update_combo():
    configs = load_configs()
    config_combo['values'] = list(configs.keys())

def on_config_select(event=None):
    selected_name = config_combo.get()
    configs = load_configs()
    if selected_name in configs:
        c = configs[selected_name]
        title_entry.delete(0, tk.END); title_entry.insert(0, c.get("sel_title", ""))
        content_entry.delete(0, tk.END); content_entry.insert(0, c.get("sel_content", ""))
        next_entry.delete(0, tk.END); next_entry.insert(0, c.get("sel_next", ""))
        remove_entry.delete(0, tk.END); remove_entry.insert(0, c.get("sel_remove", ""))
        text_remove_entry.delete(0, tk.END); text_remove_entry.insert(0, c.get("text_remove", ""))

def save_current_config():
    name = config_name_entry.get().strip()
    if not name:
        messagebox.showwarning("Lỗi", "Vui lòng nhập Tên Cấu Hình để lưu!")
        return
    configs = load_configs()
    configs[name] = {
        "sel_title": title_entry.get().strip(),
        "sel_content": content_entry.get().strip(),
        "sel_next": next_entry.get().strip(),
        "sel_remove": remove_entry.get().strip(),
        "text_remove": text_remove_entry.get().strip()
    }
    save_configs_to_file(configs)
    update_combo()
    config_combo.set(name) 
    messagebox.showinfo("Thành công", f"Đã lưu cấu hình: {name}")

def choose_directory():
    folder_selected = filedialog.askdirectory(title="Chọn thư mục lưu truyện")
    if folder_selected:
        path_entry.delete(0, tk.END)
        path_entry.insert(0, folder_selected)

def clean_filename(name):
    return re.sub(r'[\\/*?:"<>|]', "", name)

# ================= CÁC HÀM XỬ LÝ EXCEL =================
def download_sample_excel():
    file_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx", 
        filetypes=[("Excel files", "*.xlsx")],
        initialfile="Mau_Danh_Sach_Truyen.xlsx",
        title="Lưu file Excel mẫu"
    )
    if not file_path: return
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh sách truyện"
        headers = ["STT", "Link chương 1", "Tên Truyện", "Từ chương", "Số lượng tải"]
        ws.append(headers)
        data = [
            [1, "https://example.com/truyen-a/chuong-1", "Truyện A", 1, 1000],
            [2, "", "", "", ""]
        ]
        for row in data: ws.append(row)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        for cell in ws[1]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = alignment; cell.border = thin_border

        ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 30; ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 15

        wb.save(file_path)
        messagebox.showinfo("Thành công", f"Đã lưu file mẫu tại:\n{file_path}")
        os.startfile(file_path)
    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể tạo file Excel:\n{str(e)}")

def import_excel_data():
    file_path = filedialog.askopenfilename(title="Chọn file Excel", filetypes=[("Excel files", "*.xlsx *.xls")])
    if not file_path: return
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        row_idx = 2
        for i in range(10): 
            if row_idx > ws.max_row: break
            link = ws.cell(row=row_idx, column=2).value
            name = ws.cell(row=row_idx, column=3).value
            start = ws.cell(row=row_idx, column=4).value
            count = ws.cell(row=row_idx, column=5).value
            
            if link and name:
                url_entries[i].delete(0, tk.END); url_entries[i].insert(0, str(link).strip())
                name_entries[i].delete(0, tk.END); name_entries[i].insert(0, str(name).strip())
                start_entries[i].delete(0, tk.END); start_entries[i].insert(0, str(start) if start else "1")
                count_entries[i].delete(0, tk.END); count_entries[i].insert(0, str(count) if count else "1000")
            row_idx += 1
        messagebox.showinfo("Thành công", "Đã nạp danh sách truyện từ Excel thành công!")
    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể đọc file Excel. Chi tiết: {str(e)}")

# ================= CÁC HÀM ĐIỀU KHIỂN NÚT BẤM =================
def pause_bot():
    pause_event.clear()
    pause_btn.config(state=tk.DISABLED); resume_btn.config(state=tk.NORMAL)
    status_label.config(text="Trạng thái: ⏸ ĐÃ TẠM DỪNG", fg="red")

def resume_bot():
    pause_event.set()
    resume_btn.config(state=tk.DISABLED); pause_btn.config(state=tk.NORMAL)
    status_label.config(text="Trạng thái: ▶️ Đang tiếp tục chạy...", fg="green")

def reset_buttons():
    start_btn.config(state=tk.NORMAL)
    pause_btn.config(state=tk.DISABLED); resume_btn.config(state=tk.DISABLED)

# ================= HÀM CHẠY BOT CỐT LÕI =================
def run_bot(base_path, task_list, sel_title, sel_content, sel_next, sel_remove, text_remove, delay_time):
    if not base_path: base_path = os.getcwd()

    options = webdriver.EdgeOptions()
    options.add_argument('--headless') 
    options.add_argument('--disable-gpu')
    options.add_argument('--window-size=1920,1080')
    options.add_argument('user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 Edg/120.0.0.0')
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option('useAutomationExtension', False)

    driver = webdriver.Edge(options=options)
    total_tasks = len(task_list)

    text_remove_list = []
    if text_remove:
        text_remove_list = [t.strip() for t in text_remove.split(',') if t.strip()]

    try:
        for task_idx, task in enumerate(task_list, 1):
            url = task['url']
            story_name = task['name']
            num_chapters = task['count']
            current_chap = task['start']
            
            save_dir = os.path.join(base_path, clean_filename(story_name))
            if not os.path.exists(save_dir): os.makedirs(save_dir)

            status_label.config(text=f"Trạng thái: [Truyện {task_idx}/{total_tasks}] Khởi động...", fg="blue")
            driver.get(url)
            
            copied_count = 0

            while copied_count < num_chapters:
                pause_event.wait()
                status_label.config(text=f"Trạng thái: [{task_idx}/{total_tasks}] Tải '{story_name}' - Chương {current_chap}", fg="orange")
                time.sleep(delay_time)

                if sel_remove:
                    try: driver.execute_script(f"document.querySelectorAll('{sel_remove}').forEach(el => el.remove());")
                    except: pass

                title_text = ""
                if sel_title:
                    try: title_text = driver.find_elements(By.CSS_SELECTOR, sel_title)[0].text.strip()
                    except: pass

                content_text = ""
                if sel_content:
                    try: content_text = driver.find_elements(By.CSS_SELECTOR, sel_content)[0].text.strip()
                    except: pass
                
                final_content = f"{title_text}\n\n{content_text}"

                if text_remove_list and final_content:
                    for garbage_text in text_remove_list:
                        final_content = final_content.replace(garbage_text, "")
                
                final_content = final_content.strip()

                file_name = f"{save_dir}/{clean_filename(story_name)}_Chuong_{current_chap}.txt"
                with open(file_name, "w", encoding="utf-8") as f:
                    f.write(final_content)
                print(f"✅ [{story_name}] Đã lưu: {file_name}")

                copied_count += 1
                if copied_count >= num_chapters: break

                next_btn = None
                attempts = 0
                
                while attempts < 2:
                    pause_event.wait()
                    next_btn = None
                    if sel_next:
                        try: next_btn = driver.find_elements(By.CSS_SELECTOR, sel_next)[0]
                        except: pass
                    if not next_btn:
                        for link in driver.find_elements(By.TAG_NAME, "a"):
                            if "chương tiếp" in link.text.lower() or "next chapter" in link.text.lower():
                                next_btn = link; break
                    
                    if next_btn: break 
                    
                    if attempts == 0:
                        for i in range(60, 0, -1):
                            pause_event.wait()
                            if pause_event.is_set():
                                status_label.config(text=f"Trạng thái: [{task_idx}/{total_tasks}] Không thấy Next. Đợi {i}s F5...", fg="magenta")
                            time.sleep(1)
                        pause_event.wait()
                        if pause_event.is_set():
                            status_label.config(text="Trạng thái: 🔄 Đang F5 kiểm tra lại lần cuối...", fg="blue")
                            driver.refresh(); time.sleep(3) 
                    attempts += 1

                if next_btn:
                    time.sleep(delay_time / 2)
                    try:
                        n_url = next_btn.get_attribute("href")
                        if n_url and n_url != "#" and "javascript" not in n_url: driver.get(n_url)
                        else: next_btn.click()
                    except: driver.execute_script("arguments[0].click();", next_btn)
                    current_chap += 1
                else:
                    print(f"🎉 Truyện '{story_name}' đã hết. Đang nhảy sang tác vụ tiếp theo!")
                    break 

        status_label.config(text="Trạng thái: ĐÃ HOÀN THÀNH TẤT CẢ TÁC VỤ!", fg="green")
        messagebox.showinfo("Hoàn thành", "Đã xử lý xong toàn bộ danh sách truyện!")

    except Exception as e:
        status_label.config(text="Trạng thái: Bị lỗi!", fg="red")
        messagebox.showerror("Lỗi hệ thống", f"Chi tiết lỗi:\n{str(e)}")
    finally:
        driver.quit()
        reset_buttons()

def start_thread():
    base_path = path_entry.get().strip()
    sel_title = title_entry.get().strip(); sel_content = content_entry.get().strip()
    sel_next = next_entry.get().strip(); sel_remove = remove_entry.get().strip()
    text_remove = text_remove_entry.get().strip() 
    
    try: delay_time = float(delay_entry.get().strip())
    except: messagebox.showerror("Lỗi", "Độ trễ phải là số!"); return
    if not sel_next: messagebox.showerror("Lỗi", "Vui lòng nhập Selector Nút Next!"); return

    task_list = []
    for i in range(10):
        u, n = url_entries[i].get().strip(), name_entries[i].get().strip()
        if u and n:
            try: task_list.append({'url': u, 'name': n, 'start': int(start_entries[i].get().strip()), 'count': int(count_entries[i].get().strip())})
            except: messagebox.showerror("Lỗi", f"Dòng {i+1} nhập số không đúng!"); return

    if not task_list: messagebox.showerror("Lỗi", "Danh sách trống!"); return

    start_btn.config(state=tk.DISABLED); pause_btn.config(state=tk.NORMAL); resume_btn.config(state=tk.DISABLED); pause_event.set()
    threading.Thread(target=run_bot, args=(base_path, task_list, sel_title, sel_content, sel_next, sel_remove, text_remove, delay_time), daemon=True).start()

# ================= GIAO DIỆN CHÍNH (UI) =================
root = tk.Tk()
root.title(f"Auto Novel Downloader Pro Max (v{APP_VERSION})")
root.geometry("820x920") 
root.configure(padx=15, pady=10)

frame_top = tk.LabelFrame(root, text="Thiết lập Chung", padx=10, pady=10, font=("Arial", 9, "bold"))
frame_top.pack(fill="x", pady=5)
tk.Label(frame_top, text="Nơi lưu file:").grid(row=0, column=0, sticky="w")
path_entry = tk.Entry(frame_top, width=45); path_entry.grid(row=0, column=1, padx=5, pady=2)
tk.Button(frame_top, text="Chọn Thư Mục", command=choose_directory, bg="#e2e6ea").grid(row=0, column=2)
tk.Label(frame_top, text="Mẫu trang web:").grid(row=1, column=0, sticky="w", pady=5)
config_combo = ttk.Combobox(frame_top, width=42, state="readonly"); config_combo.grid(row=1, column=1, padx=5, pady=5)
config_combo.bind("<<ComboboxSelected>>", on_config_select)
tk.Label(frame_top, text="Độ trễ (s):", fg="red").grid(row=1, column=2, sticky="e", padx=(10,2))
delay_entry = tk.Entry(frame_top, width=5); delay_entry.insert(0, "1.0"); delay_entry.grid(row=1, column=3)

frame_task = tk.LabelFrame(root, text="Danh sách Truyện", padx=10, pady=10, font=("Arial", 9, "bold"))
frame_task.pack(fill="x", pady=5)

excel_frame = tk.Frame(frame_task)
excel_frame.grid(row=0, column=0, columnspan=5, pady=(0, 10), sticky="w")
tk.Button(excel_frame, text="📄 Tải File Excel Mẫu", command=download_sample_excel, bg="#28a745", fg="white", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 10))
tk.Button(excel_frame, text="📥 Nhập từ Excel", command=import_excel_data, bg="#17a2b8", fg="white", font=("Arial", 9, "bold")).pack(side="left")

tk.Label(frame_task, text="STT", font=("Arial", 8, "bold")).grid(row=1, column=0)
tk.Label(frame_task, text="Link chương bắt đầu", font=("Arial", 8, "bold")).grid(row=1, column=1, sticky="w", padx=5)
tk.Label(frame_task, text="Tên Truyện", font=("Arial", 8, "bold")).grid(row=1, column=2, sticky="w", padx=5)
tk.Label(frame_task, text="Từ C.", font=("Arial", 8, "bold")).grid(row=1, column=3)
tk.Label(frame_task, text="Số tải", font=("Arial", 8, "bold")).grid(row=1, column=4)

url_entries, name_entries, start_entries, count_entries = [], [], [], []
for i in range(10):
    tk.Label(frame_task, text=f"{i+1}.").grid(row=i+2, column=0) 
    u = tk.Entry(frame_task, width=40); u.grid(row=i+2, column=1, padx=5, pady=2); url_entries.append(u)
    n = tk.Entry(frame_task, width=28); n.grid(row=i+2, column=2, padx=5, pady=2); name_entries.append(n)
    s = tk.Entry(frame_task, width=6); s.insert(0, "1"); s.grid(row=i+2, column=3, padx=5); start_entries.append(s)
    c = tk.Entry(frame_task, width=6); c.insert(0, "1000"); c.grid(row=i+2, column=4, padx=5); count_entries.append(c)

frame_sel = tk.LabelFrame(root, text="Cấu hình Code Trang Web (Dùng chung cho cả danh sách)", padx=10, pady=10, font=("Arial", 9, "bold"))
frame_sel.pack(fill="x", pady=5)
tk.Label(frame_sel, text="Tiêu đề:").grid(row=0, column=0, sticky="w")
title_entry = tk.Entry(frame_sel, width=60); title_entry.grid(row=0, column=1, pady=2, padx=5)
tk.Label(frame_sel, text="Nội dung (*):").grid(row=1, column=0, sticky="w")
content_entry = tk.Entry(frame_sel, width=60); content_entry.grid(row=1, column=1, pady=2, padx=5)
tk.Label(frame_sel, text="Nút Next (*):").grid(row=2, column=0, sticky="w")
next_entry = tk.Entry(frame_sel, width=60); next_entry.grid(row=2, column=1, pady=2, padx=5)
tk.Label(frame_sel, text="Xóa rác (Thẻ HTML):").grid(row=3, column=0, sticky="w")
remove_entry = tk.Entry(frame_sel, width=60); remove_entry.grid(row=3, column=1, pady=2, padx=5)
tk.Label(frame_sel, text="Xóa chữ (Văn bản):", fg="#d32f2f").grid(row=4, column=0, sticky="w")
text_remove_entry = tk.Entry(frame_sel, width=60); text_remove_entry.grid(row=4, column=1, pady=2, padx=5)

frame_save_cfg = tk.Frame(frame_sel)
frame_save_cfg.grid(row=5, column=0, columnspan=2, pady=10, sticky="w")
tk.Label(frame_save_cfg, text="Tên cấu hình:").pack(side="left")
config_name_entry = tk.Entry(frame_save_cfg, width=20); config_name_entry.pack(side="left", padx=5)
tk.Button(frame_save_cfg, text="Lưu Cấu Hình", command=save_current_config, bg="#ffc107", font=("Arial", 8, "bold")).pack(side="left")

ctrl_frame = tk.Frame(root)
ctrl_frame.pack(fill="x", pady=10)
start_btn = tk.Button(ctrl_frame, text="🚀 CHẠY DANH SÁCH", bg="#007bff", fg="white", font=("Arial", 11, "bold"), height=2, command=start_thread); start_btn.pack(side="left", fill="x", expand=True, padx=2)
pause_btn = tk.Button(ctrl_frame, text="⏸ TẠM DỪNG", bg="#ffc107", font=("Arial", 11, "bold"), height=2, state=tk.DISABLED, command=pause_bot); pause_btn.pack(side="left", fill="x", expand=True, padx=2)
resume_btn = tk.Button(ctrl_frame, text="▶️ TIẾP TỤC", bg="#28a745", fg="white", font=("Arial", 11, "bold"), height=2, state=tk.DISABLED, command=resume_bot); resume_btn.pack(side="left", fill="x", expand=True, padx=2)

status_label = tk.Label(root, text="Trạng thái: Sẵn sàng", font=("Arial", 10, "italic"), fg="gray"); status_label.pack()

update_combo()
root.mainloop()
